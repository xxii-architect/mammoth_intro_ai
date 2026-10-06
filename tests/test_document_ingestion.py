import asyncio
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi import UploadFile
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import api_server
from mammoth_os.documents import (
    DocumentError,
    DocumentLibrary,
    UploadBodyLimit,
    capabilities,
    extract_document,
    retrieve_sections,
)


def make_pdf(path, pages=120, blank=False):
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if not blank:
            text = ("Routine teaching background with foundational concepts. " * 12
                    if index < pages - 1 else "Zebrafish regeneration uses cardiac progenitor cells. This is the final chapter.")
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(content.flate_encode())
    writer.write(path)


@pytest.fixture
def library(tmp_path):
    return DocumentLibrary(tmp_path / "documents", {"chat": tmp_path / "old-chat", "atlas": tmp_path / "old-atlas"})


def upload(library, data, filename="notes.txt", scope="atlas"):
    return asyncio.run(library.upload(UploadFile(filename=filename, file=io.BytesIO(data)), scope))


def test_default_policy_and_success_over_old_four_mb_limit(library):
    policy = capabilities()
    assert policy["max_file_bytes"] == 50 * 1024 * 1024
    result = upload(library, b"Large education text.\n" * 210_000)
    assert result["size"] > 4 * 1024 * 1024
    assert result["processing_status"] == "partial"
    assert 8000 < result["extracted_chars"] <= 2_000_000
    assert any("limit" in warning for warning in result["warnings"])
    assert "path" not in result


def test_actual_fifty_mib_file_is_accepted_and_one_byte_more_is_rejected(library):
    maximum = capabilities()["max_file_bytes"]
    result = upload(library, b"a" * maximum)
    assert result["size"] == 50 * 1024 * 1024
    assert result["processing_status"] == "partial"
    assert result["extracted_chars"] == 2_000_000
    with pytest.raises(DocumentError) as error:
        upload(library, b"a" * (maximum + 1))
    assert error.value.status_code == 413
    assert library.usage()["files"] == 1


def test_late_compressed_pdf_page_is_retrieved_and_located(library, tmp_path):
    path = tmp_path / "textbook.pdf"
    make_pdf(path)
    result = upload(library, path.read_bytes(), "textbook.pdf")
    assert result["processing_status"] == "ready"
    assert result["extracted_chars"] > 8000
    entry = library.list("atlas")[0]
    sections = library.sections(entry, "Explain zebrafish cardiac regeneration")
    assert sections[0]["location"] == "page 120"
    assert "cardiac progenitor" in sections[0]["text"]
    assert library.sections(entry, "page 120")[0]["location"] == "page 120"


def test_real_docx_pptx_and_xlsx_readers(library, tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    document = Document()
    document.add_heading("Thermodynamics", 1)
    document.add_paragraph("Entropy quantifies energy dispersal.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Heat capacity table"
    docx = tmp_path / "book.docx"
    document.save(docx)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Photosynthesis"
    slide.placeholders[1].text = "Chlorophyll absorbs light."
    pptx = tmp_path / "lecture.pptx"
    presentation.save(pptx)
    workbook = Workbook()
    workbook.active.title = "Measurements"
    workbook.active.append(["Temperature", "298 Kelvin"])
    xlsx = tmp_path / "lab.xlsx"
    workbook.save(xlsx)
    for path, query, expected, location in [
        (docx, "entropy", "energy dispersal", "block"),
        (pptx, "chlorophyll", "absorbs light", "slide 1"),
        (xlsx, "temperature", "298 Kelvin", "sheet Measurements"),
    ]:
        upload(library, path.read_bytes(), path.name)
        section = library.sections(library.list("atlas")[0], query)[0]
        assert expected in section["text"]
        assert section["location"].startswith(location)


def test_html_and_utf16_keep_readable_text_without_active_content(library):
    result = upload(library, b"<script>secret()</script><h1>Geometry</h1><p>Triangles have three sides.</p>", "notes.html")
    assert "secret" not in result["text_preview"]
    assert "Triangles" in result["text_preview"]
    assert "Geometry" in result["text_preview"]
    result = upload(library, "Unicode learning: caf\u00e9".encode("utf-16"))
    assert "caf\u00e9" in result["text_preview"]


@pytest.mark.parametrize("filename,data", [("program.exe", b"binary"), ("broken.pdf", b"not a pdf"), ("binary.txt", b"a\x00b"), ("empty.txt", b"")])
def test_invalid_inputs_fail_without_success_records(library, filename, data):
    with pytest.raises(DocumentError):
        upload(library, data, filename)
    assert library.list("atlas") == []
    assert not list(library.folder.glob("*.upload"))
    assert not list(library.folder.glob("*.chunks.json"))


def test_oversize_is_bounded_and_cleaned(library, monkeypatch):
    monkeypatch.setenv("MAMMOTH_UPLOAD_MAX_BYTES", "8")
    with pytest.raises(DocumentError) as caught:
        upload(library, b"123456789")
    assert caught.value.status_code == 413
    assert not list(library.folder.glob("*.upload"))


def test_quota_shared_by_chat_atlas_and_concurrent_uploads(library, monkeypatch):
    monkeypatch.setenv("MAMMOTH_UPLOAD_STORAGE_BYTES", "10")
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = [pool.submit(upload, library, b"123456", "notes.txt", scope) for scope in ("chat", "atlas")]
        successes = failures = 0
        for result in results:
            try:
                result.result()
                successes += 1
            except DocumentError:
                failures += 1
    assert (successes, failures) == (1, 1)
    assert library.usage() == {"storage_bytes": 6, "files": 1}


def test_ocr_needed_and_encrypted_pdf_are_not_claimed_ready(library, tmp_path):
    path = tmp_path / "scan.pdf"
    make_pdf(path, pages=1, blank=True)
    result = upload(library, path.read_bytes(), path.name)
    assert result["processing_status"] == "needs_ocr"
    assert result["chunk_count"] == 0
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(DocumentError, match="Encrypted"):
        upload(library, path.read_bytes(), path.name)


def test_archive_bomb_and_xml_entities_rejected(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * 2_000_000)
    with pytest.raises(DocumentError):
        extract_document(path, path.name)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///private">]><foo>&xxe;</foo>')
    with pytest.raises(DocumentError):
        extract_document(path, path.name)


def test_legacy_migration_retains_records_but_never_claims_full_extraction(library):
    folder = library.legacy_folders["atlas"]
    folder.mkdir()
    (folder / "_index.json").write_text(json.dumps([{"file_id": "old", "name": "notes.md", "size": 20, "text_preview": "Legacy teaching text"}]))
    entries = library.list("atlas")
    assert entries[0]["processing_status"] == "legacy_preview"
    assert library.sections(entries[0], "teaching")[0]["location"] == "legacy preview"
    assert len(library.list("atlas")) == 1


def test_corrupt_legacy_index_is_explicit_not_empty_success(library):
    folder = library.legacy_folders["chat"]
    folder.mkdir()
    (folder / "_index.json").write_text("{broken")
    with pytest.raises(DocumentError, match="unreadable"):
        library.list("chat")


def test_routes_private_metadata_and_full_document_context(monkeypatch, tmp_path):
    monkeypatch.setattr(api_server, "MAMMOTH_DIR", tmp_path)
    monkeypatch.setattr(api_server, "USER_UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setattr(api_server, "ATLAS_FILES_DIR", tmp_path / "atlas")
    (tmp_path / "uploads").mkdir()
    (tmp_path / "atlas").mkdir()
    monkeypatch.setattr(api_server, "_current_request_user_id", lambda: "user-a")
    path = tmp_path / "textbook.pdf"
    make_pdf(path)
    result = asyncio.run(api_server.upload_atlas_file(UploadFile(filename=path.name, file=io.BytesIO(path.read_bytes())), tag="textbook"))
    context = api_server._collect_attached_atlas_material_context("user-a", [result["file_id"]], "zebrafish regeneration")
    assert "page 120" in context["materials"][0]["excerpt"]
    assert "cardiac progenitor" in context["materials"][0]["excerpt"]
    public = asyncio.run(api_server.list_atlas_files())
    assert "path" not in public["files"][0]
    assert "chunks_path" not in public["files"][0]
    content = asyncio.run(api_server.get_atlas_file_content(result["file_id"], "zebrafish"))
    assert content["sections"][0]["location"] == "page 120"
    monkeypatch.setattr(api_server, "_current_request_user_id", lambda: "user-b")
    assert asyncio.run(api_server.list_atlas_files())["files"] == []
    assert asyncio.run(api_server.get_atlas_file_content(result["file_id"])).status_code == 404
    assert asyncio.run(api_server.delete_atlas_file(result["file_id"])).status_code == 404
    assert api_server._collect_attached_atlas_material_context("user-b", [result["file_id"]], "zebrafish")["count"] == 0


def test_delete_and_retag_are_persistent_and_scope_checked(library):
    result = upload(library, b"Private lecture text.")
    with pytest.raises(DocumentError):
        library.delete("chat", result["file_id"])
    library.retag("atlas", result["file_id"], "textbook")
    assert library.list("atlas")[0]["tag"] == "textbook"
    library.delete("atlas", result["file_id"])
    assert library.usage()["files"] == 0
    assert not list(library.folder.glob("*.chunks.json"))


def test_retrieval_reports_opening_fallback_not_semantic_certainty():
    assert retrieve_sections([{"location": "page 1", "text": "Some unrelated text"}], "zebrafish")[0]["retrieval"] == "opening_sections"


def test_multipart_request_limit_covers_declared_and_chunked_bodies(monkeypatch):
    from fastapi import FastAPI, File, UploadFile
    from fastapi.testclient import TestClient
    monkeypatch.setenv("MAMMOTH_UPLOAD_MAX_BYTES", "16")
    app = FastAPI()
    app.add_middleware(UploadBodyLimit)
    file_parameter = File(...)

    @app.post("/api/atlas/files/upload")
    async def receive_file(file: UploadFile = file_parameter):
        return {"size": len(await file.read())}

    with TestClient(app) as client:
        result = client.post("/api/atlas/files/upload", content=b"tiny", headers={"content-length": str(2 * 1024 * 1024)})
        assert result.status_code == 413
        def chunks():
            for _ in range(5):
                yield b"x" * (256 * 1024)
        result = client.post("/api/atlas/files/upload", content=chunks(), headers={"content-type": "multipart/form-data; boundary=test"})
        assert result.status_code == 413
        from mammoth_os import documents
        assert documents._REQUEST_UPLOAD_SLOTS.acquire(blocking=False)
        assert documents._REQUEST_UPLOAD_SLOTS.acquire(blocking=False)
        try:
            assert client.post("/api/atlas/files/upload", content=b"tiny").status_code == 429
        finally:
            documents._REQUEST_UPLOAD_SLOTS.release()
            documents._REQUEST_UPLOAD_SLOTS.release()


def test_parser_timeout_cleanup_is_explicit(monkeypatch, tmp_path):
    import subprocess

    from mammoth_os import documents
    path = tmp_path / "staged.upload"
    path.write_bytes(b"readable text")
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("reader", 60)
    monkeypatch.setattr(documents.subprocess, "run", timeout)
    with pytest.raises(DocumentError, match="60 seconds"):
        documents._extract_bounded(path, "notes.txt")
    assert not path.with_suffix(".extraction.json").exists()


def test_chat_and_streaming_context_use_late_sections(monkeypatch, tmp_path):
    monkeypatch.setattr(api_server, "MAMMOTH_DIR", tmp_path)
    monkeypatch.setattr(api_server, "USER_UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setattr(api_server, "ATLAS_FILES_DIR", tmp_path / "atlas")
    (tmp_path / "uploads").mkdir()
    (tmp_path / "atlas").mkdir()
    monkeypatch.setattr(api_server, "_current_request_user_id", lambda: "reader")
    content = ("Ordinary foundation material.\n" * 1000 + "Late chapter: zebrafish cardiac progenitor regeneration.\n").encode()
    result = asyncio.run(api_server.upload_chat_file(UploadFile(filename="book.txt", file=io.BytesIO(content))))
    context = api_server._attached_chat_document_context("reader", [result["file_id"]], "zebrafish cardiac regeneration")
    assert "Late chapter" in context
    assert "lines 1001" in context
    assert api_server._attached_chat_document_context("different-reader", [result["file_id"]], "zebrafish") == ""


def test_exact_file_limit_is_accepted_and_busy_slots_are_retryable(library, monkeypatch):
    from mammoth_os import documents
    monkeypatch.setenv("MAMMOTH_UPLOAD_MAX_BYTES", "8")
    assert upload(library, b"12345678")["size"] == 8
    assert documents._UPLOAD_SLOTS.acquire(blocking=False)
    assert documents._UPLOAD_SLOTS.acquire(blocking=False)
    try:
        with pytest.raises(DocumentError) as error:
            upload(library, b"12345678")
        assert error.value.status_code == 429
    finally:
        documents._UPLOAD_SLOTS.release()
        documents._UPLOAD_SLOTS.release()


def test_atlas_model_receives_late_page_and_returns_location(monkeypatch, tmp_path):
    from mammoth_os import llm_client
    library = DocumentLibrary(tmp_path / "private", {"chat": tmp_path / "old-chat", "atlas": tmp_path / "old-atlas"})
    path = tmp_path / "textbook.pdf"
    make_pdf(path)
    uploaded = upload(library, path.read_bytes(), path.name)
    monkeypatch.setattr(api_server, "_document_library", lambda user_id: library)
    monkeypatch.setattr(api_server, "_current_request_user_id", lambda: "reader")
    state = {"current_exercise": {}, "current_lesson": {}, "lesson_plan": {}, "resume_packet": {}}
    monkeypatch.setattr(api_server, "_load_atlas_state", lambda: state)
    monkeypatch.setattr(api_server, "_save_atlas_state", lambda state: None)
    monkeypatch.setattr(api_server, "_hydrate_learner_state", lambda *args, **kwargs: {})
    monkeypatch.setattr(api_server, "_sync_resume_packet", lambda *args, **kwargs: None)
    class Client:
        model = "test-only"
        async def generate(self, prompt, **kwargs):
            assert "page 120" in prompt
            assert "cardiac progenitor cells" in prompt
            return "Cardiac progenitor cells contribute to regeneration (textbook.pdf, page 120)."
    monkeypatch.setattr(llm_client, "get_llm_client", lambda *args, **kwargs: Client())
    result = asyncio.run(api_server.atlas_chat({
        "message": "Explain zebrafish cardiac regeneration.", "mode": "assistant",
        "attached_material_ids": [uploaded["file_id"]],
    }))
    assert "page 120" in result["reply"]
    assert "page 120" in result["evidence_items"][0]["summary"]
