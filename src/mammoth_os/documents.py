"""Private, bounded document ingestion and deterministic source-section retrieval."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sqlite3
import subprocess
import sys
import threading
import uuid
import zipfile
from contextlib import closing
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".py", ".js", ".jsx", ".ts", ".tsx", ".json",
    ".toml", ".yaml", ".yml", ".css", ".sh", ".sql", ".rst", ".log",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"}
MAX_EXTRACTED_CHARS = 2_000_000
MAX_UNITS = 2000
MAX_ARCHIVE_BYTES = 200 * 1024 * 1024
_UPLOAD_SLOTS = threading.BoundedSemaphore(2)
_REQUEST_UPLOAD_SLOTS = threading.BoundedSemaphore(2)


class DocumentError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


def capabilities() -> dict:
    policy = {
        "contract": "mammoth.documents.v1",
        "extensions": sorted(TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS),
        "max_file_bytes": int(os.environ.get("MAMMOTH_UPLOAD_MAX_BYTES", "52428800")),
        "max_storage_bytes": int(os.environ.get("MAMMOTH_UPLOAD_STORAGE_BYTES", "524288000")),
        "max_files": int(os.environ.get("MAMMOTH_UPLOAD_MAX_FILES", "200")),
        "max_extracted_chars": MAX_EXTRACTED_CHARS,
        "max_units": MAX_UNITS,
        "ocr": "not_enabled",
        "media_transcription": "not_enabled",
    }
    if any(policy[key] <= 0 for key in ("max_file_bytes", "max_storage_bytes", "max_files")):
        raise DocumentError("Upload configuration must use positive limits.", 503)
    return policy


class UploadBodyLimit:
    """Bound multipart bodies before Starlette spools them to disk."""

    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http" or scope.get("method") != "POST" or scope.get("path") not in {
            "/api/atlas/files/upload", "/api/mammoth/files/upload",
        }:
            return await self.app(scope, receive, send)
        from starlette.responses import JSONResponse
        limit = capabilities()["max_file_bytes"] + 1024 * 1024
        headers = dict(scope.get("headers", []))
        length = headers.get(b"content-length")
        if length:
            try:
                declared = int(length)
            except ValueError:
                return await JSONResponse({"status": "error", "error": "Invalid upload content length."}, status_code=400)(scope, receive, send)
            if declared > limit:
                return await JSONResponse({"status": "error", "error": "Upload request exceeds the file size limit."}, status_code=413)(scope, receive, send)
        total = 0
        exceeded = False

        async def bounded_receive():
            nonlocal total, exceeded
            event = await receive()
            total += len(event.get("body", b""))
            if total > limit:
                exceeded = True
                from starlette.exceptions import HTTPException
                raise HTTPException(413, "Upload request exceeds the file size limit.")
            return event

        async def bounded_send(event):
            if exceeded and event["type"] == "http.response.start":
                event = {**event, "status": 413}
            await send(event)

        if not _REQUEST_UPLOAD_SLOTS.acquire(blocking=False):
            return await JSONResponse({"status": "error", "error": "Document uploads are busy. Retry when another upload finishes."}, status_code=429)(scope, receive, send)
        try:
            await self.app(scope, bounded_receive, bounded_send)
        finally:
            _REQUEST_UPLOAD_SLOTS.release()


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag in {"p", "div", "br", "h1", "h2", "h3", "li", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _check_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 10_000 or sum(item.file_size for item in members) > MAX_ARCHIVE_BYTES:
            raise DocumentError("Document expands beyond the safe processing limit.")
        if any(item.flag_bits & 1 for item in members):
            raise DocumentError("Encrypted documents are not supported.")
        if any(item.file_size > 1024 * 1024 and item.file_size > max(item.compress_size, 1) * 200 for item in members):
            raise DocumentError("Document compression ratio exceeds the safe processing limit.")
        if any("vbaproject" in item.filename.lower() for item in members):
            raise DocumentError("Macro-enabled documents are not supported.")
        from defusedxml.ElementTree import fromstring
        for item in members:
            if item.filename.endswith((".xml", ".rels")):
                fromstring(archive.read(item))


def _units(path: Path, extension: str):
    if extension in TEXT_EXTENSIONS | {".html", ".htm"}:
        raw = path.read_bytes()
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            text = raw.decode(encoding)
        except UnicodeError as exc:
            raise DocumentError("Text must use UTF-8 or BOM-marked UTF-16 encoding.") from exc
        if "\x00" in text:
            raise DocumentError("This file contains binary data, not readable text.")
        if extension in {".html", ".htm"}:
            parser = _HTMLText()
            parser.feed(text)
            text = "".join(parser.parts)
        lines = text.splitlines(keepends=True)
        for start in range(0, len(lines), 40):
            yield f"lines {start + 1}-{min(start + 40, len(lines))}", "".join(lines[start:start + 40])
    elif extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentError("Encrypted PDFs are not supported. Upload an unlocked copy.")
        for index, page in enumerate(reader.pages):
            if index >= MAX_UNITS:
                yield "processing limit", ""
                break
            yield f"page {index + 1}", page.extract_text() or ""
    elif extension == ".docx":
        from docx import Document
        document = Document(path)
        # Preserve paragraph/table order, including headings, rather than inventing pages.
        from docx.oxml.ns import qn
        for index, element in enumerate(document.element.body):
            if element.tag not in {qn("w:p"), qn("w:tbl")}:
                continue
            text = " ".join(node.text or "" for node in element.iter(qn("w:t")))
            yield f"block {index + 1}", text
    elif extension == ".pptx":
        from pptx import Presentation
        for index, slide in enumerate(Presentation(path).slides):
            text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text.append(shape.text)
                if shape.has_table:
                    text.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            yield f"slide {index + 1}", "\n".join(text)
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    rows = []
                    for index, row in enumerate(sheet.iter_rows(values_only=True)):
                        rows.append(" | ".join(str(value) if value is not None else "" for value in row))
                        if len(rows) == 40:
                            yield f"sheet {sheet.title}, rows {index - 38}-{index + 1}", "\n".join(rows)
                            rows = []
                    if rows:
                        yield f"sheet {sheet.title}, rows {index - len(rows) + 2}-{index + 1}", "\n".join(rows)
            finally:
                book.close()


def extract_document(path: Path, filename: str) -> dict:
    extension = Path(filename).suffix.lower()
    warnings = []
    if extension == ".xlsx":
        warnings.append("Spreadsheet formulas are not recalculated; only saved cached values are available.")
    if extension in {".pptx", ".docx"}:
        warnings.append("Embedded images are not transcribed. Document blocks or slide text are indexed, not rendered pages.")
    chunks = []
    total = empty = count = 0
    try:
        if extension in {".docx", ".pptx", ".xlsx"}:
            _check_archive(path)
        for count, (location, text) in enumerate(_units(path, extension), start=1):
            if count > MAX_UNITS or location == "processing limit" or total >= MAX_EXTRACTED_CHARS:
                warnings.append("Processing limit reached; only part of this document was indexed.")
                break
            text = text.strip()
            if not text:
                empty += 1
                continue
            remaining = MAX_EXTRACTED_CHARS - total
            if len(text) > remaining:
                warnings.append("Extracted text limit reached; only part of this document was indexed.")
            text = text[:remaining]
            total += len(text)
            for offset in range(0, len(text), 1200):
                chunks.append({"location": location, "text": text[offset:offset + 1400]})
    except DocumentError:
        raise
    except ImportError as exc:
        raise DocumentError("A document reader is unavailable. Contact the operator to restore server dependencies.", 503) from exc
    except Exception as exc:
        raise DocumentError("Document could not be parsed. It may be damaged or use an unsupported encoding.") from exc
    if empty:
        warnings.append(f"{empty} sections had no extractable text; images and scans require OCR, which is not enabled.")
    if not chunks:
        status = "needs_ocr" if extension == ".pdf" else "empty"
        warnings.append("No usable text was extracted. This file will not provide model context.")
    else:
        status = "partial" if warnings else "ready"
    return {"status": status, "warnings": warnings, "chunks": chunks, "extracted_chars": total, "units_processed": min(count, MAX_UNITS)}


def retrieve_sections(chunks: list[dict], query: str, *, limit: int = 4) -> list[dict]:
    words = set(re.findall(r"\w{3,}", query.casefold())) - {
        "the", "and", "this", "that", "with", "from", "what", "please", "file", "document",
    }
    locations = re.findall(r"\b(page|slide|block)\s+(\d+)\b", query.casefold())
    ranked = []
    for index, chunk in enumerate(chunks):
        tokens = set(re.findall(r"\w{3,}", chunk["text"].casefold()))
        score = len(words & tokens)
        score += 100 * sum(f"{kind} {number}" == chunk["location"].casefold() for kind, number in locations)
        ranked.append((score, index, chunk))
    ranked.sort(key=lambda row: (-row[0], row[1]))
    selected = [row[2] for row in ranked[:limit]]
    return [{**chunk, "retrieval": "lexical" if ranked and ranked[0][0] else "opening_sections"} for chunk in selected]


def _extract_bounded(path: Path, filename: str) -> dict:
    output = path.with_suffix(".extraction.json")
    try:
        subprocess.run(
            [sys.executable, "-m", "mammoth_os.documents", "--extract", str(path), filename, str(output)],
            check=True, timeout=60, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
            env={**os.environ, "PYTHONPATH": str(Path(__file__).resolve().parent.parent)},
        )
        result = json.loads(output.read_text(encoding="utf-8"))
        if "error" in result:
            raise DocumentError(result["error"], result["status_code"])
        return result
    except subprocess.TimeoutExpired as exc:
        raise DocumentError("Document processing exceeded 60 seconds. Split the document into smaller files.") from exc
    except subprocess.CalledProcessError as exc:
        raise DocumentError("Document reader stopped unexpectedly. Retry with a smaller or repaired document.") from exc
    finally:
        output.unlink(missing_ok=True)


class DocumentLibrary:
    """SQLite metadata/quota transactions; raw files and chunks never leave user scope."""

    def __init__(self, folder: Path, legacy_folders: dict[str, Path]):
        self.folder = folder
        self.legacy_folders = legacy_folders
        folder.mkdir(parents=True, exist_ok=True)

    def _connect(self):
        connection = sqlite3.connect(self.folder / "library.sqlite3", timeout=30)
        connection.execute("CREATE TABLE IF NOT EXISTS files (file_id TEXT PRIMARY KEY, scope TEXT NOT NULL, payload TEXT NOT NULL)")
        connection.execute("CREATE TABLE IF NOT EXISTS migrations (scope TEXT PRIMARY KEY)")
        try:
            for scope, folder in self.legacy_folders.items():
                if connection.execute("SELECT 1 FROM migrations WHERE scope=?", (scope,)).fetchone():
                    continue
                index = folder / "_index.json"
                if index.exists():
                    try:
                        entries = json.loads(index.read_text(encoding="utf-8"))
                    except json.JSONDecodeError as exc:
                        raise DocumentError("Stored upload index is unreadable; operator repair is required.", 500) from exc
                    if not isinstance(entries, list):
                        raise DocumentError("Stored upload index is invalid; operator repair is required.", 500)
                    for entry in entries:
                        entry = {**entry, "scope": scope, "processing_status": "legacy_preview",
                                 "warnings": ["Legacy preview only. Re-upload to index the complete document."]}
                        connection.execute("INSERT OR IGNORE INTO files VALUES (?,?,?)", (entry["file_id"], scope, json.dumps(entry)))
                connection.execute("INSERT OR IGNORE INTO migrations VALUES (?)", (scope,))
            connection.commit()
        except Exception:
            connection.close()
            raise
        return connection

    def list(self, scope: str) -> list[dict]:
        with closing(self._connect()) as connection, connection:
            rows = connection.execute("SELECT payload FROM files WHERE scope=? ORDER BY rowid DESC", (scope,)).fetchall()
        return [json.loads(row[0]) for row in rows]

    def usage(self) -> dict:
        with closing(self._connect()) as connection, connection:
            entries = [json.loads(row[0]) for row in connection.execute("SELECT payload FROM files")]
        return {"storage_bytes": sum(entry["size"] for entry in entries), "files": len(entries)}

    @staticmethod
    def public(entry: dict) -> dict:
        return {key: value for key, value in entry.items() if key not in {"path", "chunks_path", "text_preview"}}

    async def upload(self, upload, scope: str, tag: str = "other") -> dict:
        name = str(upload.filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()[:240]
        extension = Path(name).suffix.lower()
        policy = capabilities()
        if not name or extension not in policy["extensions"]:
            raise DocumentError("Unsupported file type. Consult the upload format list.", 415)
        usage = await asyncio.to_thread(self.usage)
        if usage["files"] >= policy["max_files"] or usage["storage_bytes"] >= policy["max_storage_bytes"]:
            raise DocumentError("Private upload storage quota reached. Delete unused files before uploading.", 413)
        if not _UPLOAD_SLOTS.acquire(blocking=False):
            raise DocumentError("Document processing is busy. Retry when another upload finishes.", 429)
        file_id = f"{'atlas' if scope == 'atlas' else 'file'}-{uuid.uuid4().hex[:16]}"
        size = 0
        staging = self.folder / f"{file_id}.upload"
        try:
            with staging.open("xb") as target:
                while block := await upload.read(256 * 1024):
                    size += len(block)
                    if size > policy["max_file_bytes"]:
                        raise DocumentError(f"File exceeds the {policy['max_file_bytes'] // (1024 * 1024)} MB upload limit.", 413)
                    target.write(block)
            if not size:
                raise DocumentError("Empty files cannot be uploaded.")
            # Complete processing before acknowledging readiness; never run a parser on the event loop.
            worker = asyncio.create_task(asyncio.to_thread(self._ingest, staging, file_id, name, size, scope, tag, policy))
            try:
                return await asyncio.shield(worker)
            except asyncio.CancelledError:
                # The parser cannot be killed safely mid-read; wait before removing its input.
                await worker
                raise
        finally:
            try:
                staging.unlink(missing_ok=True)
            finally:
                _UPLOAD_SLOTS.release()

    def _ingest(self, staging, file_id, name, size, scope, tag, policy):
        extracted = _extract_bounded(staging, name)
        final = self.folder / f"{file_id}{Path(name).suffix.lower()}"
        chunks_path = self.folder / f"{file_id}.chunks.json"
        entry = {
            "file_id": file_id, "name": name, "ext": Path(name).suffix.lower(), "size": size,
            "tag": tag, "scope": scope, "created_at": datetime.now(timezone.utc).isoformat(),
            "path": str(final), "chunks_path": str(chunks_path),
            "processing_status": extracted["status"], "warnings": extracted["warnings"],
            "extracted_chars": extracted["extracted_chars"], "units_processed": extracted["units_processed"],
            "chunk_count": len(extracted["chunks"]),
        }
        connection = self._connect()
        try:
            connection.execute("BEGIN IMMEDIATE")
            entries = [json.loads(row[0]) for row in connection.execute("SELECT payload FROM files")]
            if len(entries) >= policy["max_files"] or sum(item["size"] for item in entries) + size > policy["max_storage_bytes"]:
                raise DocumentError("Private upload storage quota reached. Delete unused files before uploading.", 413)
            staging.replace(final)
            chunks_path.write_text(json.dumps(extracted["chunks"], ensure_ascii=False), encoding="utf-8")
            connection.execute("INSERT INTO files VALUES (?,?,?)", (file_id, scope, json.dumps(entry)))
            connection.commit()
        except BaseException:
            connection.rollback()
            final.unlink(missing_ok=True)
            chunks_path.unlink(missing_ok=True)
            raise
        finally:
            connection.close()
        preview = "\n\n".join(chunk["text"] for chunk in extracted["chunks"][:2])[:1200]
        return {"status": "ok", **self.public(entry), "text_preview": preview}

    def sections(self, entry: dict, query: str) -> list[dict]:
        if entry.get("chunks_path"):
            path = self.folder / f"{entry['file_id']}.chunks.json"
            if not path.exists():
                raise DocumentError("Document index is missing. Re-upload the file or contact the operator.", 500)
            try:
                chunks = json.loads(path.read_text(encoding="utf-8"))
            except json.JSONDecodeError as exc:
                raise DocumentError("Document index is unreadable. Re-upload or contact the operator.", 500) from exc
            if not isinstance(chunks, list) or not all(isinstance(chunk, dict) and isinstance(chunk.get("text"), str) and isinstance(chunk.get("location"), str) for chunk in chunks):
                raise DocumentError("Document index has an invalid shape. Re-upload or contact the operator.", 500)
        else:
            chunks = [{"location": "legacy preview", "text": str(entry.get("text_preview") or "")[:8000]}]
        return retrieve_sections(chunks, query)

    def delete(self, scope: str, file_id: str) -> dict:
        connection = self._connect()
        try:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute("SELECT payload FROM files WHERE scope=? AND file_id=?", (scope, file_id)).fetchone()
            if not row:
                raise DocumentError("File not found.", 404)
            entry = json.loads(row[0])
            allowed = [self.folder.resolve(), self.legacy_folders[scope].resolve()]
            for key in ("path", "chunks_path"):
                if entry.get(key):
                    path = Path(entry[key]).resolve()
                    if path.parent not in allowed:
                        raise DocumentError("Stored file location is invalid; operator repair is required.", 500)
                    path.unlink(missing_ok=True)
            connection.execute("DELETE FROM files WHERE file_id=?", (file_id,))
            connection.commit()
        finally:
            connection.close()
        return {"status": "ok", "deleted": file_id}

    def retag(self, scope: str, file_id: str, tag: str) -> dict:
        with closing(self._connect()) as connection, connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute("SELECT payload FROM files WHERE scope=? AND file_id=?", (scope, file_id)).fetchone()
            if not row:
                raise DocumentError("File not found.", 404)
            entry = json.loads(row[0])
            entry["tag"] = tag
            connection.execute("UPDATE files SET payload=? WHERE file_id=?", (json.dumps(entry), file_id))
        return {"status": "ok", "file_id": file_id, "tag": tag}


if __name__ == "__main__":
    if len(sys.argv) != 5 or sys.argv[1] != "--extract":
        raise SystemExit("Use the document API, not this internal reader entry point.")
    if os.name == "posix":
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (1024 * 1024 * 1024, 1024 * 1024 * 1024))
    try:
        result = extract_document(Path(sys.argv[2]), sys.argv[3])
    except DocumentError as exc:
        result = {"error": str(exc), "status_code": exc.status_code}
    Path(sys.argv[4]).write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
